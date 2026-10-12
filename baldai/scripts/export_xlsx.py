"""model.json -> detales.xlsx (pjovimo lapas pjaustymo centrui).

    python scripts/export_xlsx.py <dir>/model.json [--out <dir>/detales.xlsx]

Lapai: Suvestine (formules), po viena lapa kiekvienai medziagai, Furnitura, Standartas.
Po irasymo paleisti xlsx skill'o recalc.py (formuliu reiksmes) - zr. SKILL.md.
"""
import argparse
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument("model")
ap.add_argument("--out")
args = ap.parse_args()
MODEL = Path(args.model)
model = json.loads(MODEL.read_text(encoding="utf-8"))

F = "Arial"
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
h1 = Font(name=F, size=14, bold=True)
grey = Font(name=F, size=9, color="666666")
head_fill = PatternFill("solid", fgColor="DDE3EA")
input_fill = PatternFill("solid", fgColor="FFF8DC")
thin = Side(style="thin", color="B7BEC7")
border = Border(top=thin, bottom=thin, left=thin, right=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

wb = Workbook()
wb.remove(wb.active)


def header_row(ws, row, labels, widths=None):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font, cell.fill, cell.border, cell.alignment = bold, head_fill, border, center
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def cut_sheet(grp):
    name = {"korpusas": f'LMDP{grp["storis"]} korpusai', "fasadas:apatiniai": f'LMDP{grp["storis"]} fas. apatiniai',
            "fasadas:virsutiniai": f'LMDP{grp["storis"]} fas. viršutiniai', "fasadas": f'LMDP{grp["storis"]} fasadai',
            "stalcius": f'LMDP{grp["storis"]} stalčiai', "nugarele": f'HDF{grp["storis"]} nugarėlės',
            "stalvirsis": "Stalviršis", "hpl": "HPL siena"}.get(grp["grupe"], grp["grupe"][:31])
    ws = wb.create_sheet(name)
    ws["A1"] = f'{model["projektas"]["pavadinimas"]} — {grp["pavadinimas"]}'
    ws["A1"].font = h1
    ws["A2"] = f'Dekoras: {grp["dekoras"]} ({grp["dekoro_pavadinimas"]}), storis {grp["storis"]} mm'
    ws["A2"].font = bold
    ws["A3"] = ("Matmenys galutiniai, su kantu. L (ilgis) — išilgai rašto. "
                "Kantai: L1/L2 — ilgio kraštinės, W1/W2 — pločio kraštinės, reikšmė = ABS storis mm, tuščia = be kanto. "
                "Raštas: V — vertikalus, H — horizontalus, – be rašto.")
    ws["A3"].font = grey
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:N3")
    ws.row_dimensions[3].height = 30
    labels = ["Nr.", "Detalė", "Žymėjimas", "L, mm", "W, mm", "Vnt.", "Storis", "Dekoras",
              "Kantas L1", "Kantas L2", "Kantas W1", "Kantas W2", "Raštas", "Plotas, m²"]
    widths = [5, 24, 22, 9, 9, 6, 7, 12, 9, 9, 9, 9, 7, 11]
    header_row(ws, 5, labels, widths)
    r = 6
    first = r
    for e in grp["eilutes"]:
        vals = [e["nr"], e["detale"], e["zymejimas"], e["L"], e["W"], e["vnt"], grp["storis"], grp["dekoras"],
                *[k if k else None for k in e["kantai"]], e["rastas"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = font, border
            cell.alignment = left if c in (2, 3) else center
        m2 = ws.cell(row=r, column=14, value=f"=D{r}*E{r}*F{r}/1000000")
        m2.font, m2.border, m2.alignment, m2.number_format = font, border, center, "0.000"
        r += 1
    last = r - 1
    ws.cell(row=r, column=3, value="Iš viso").font = bold
    tot = ws.cell(row=r, column=6, value=f"=SUM(F{first}:F{last})")
    tot.font, tot.alignment = bold, center
    m2t = ws.cell(row=r, column=14, value=f"=SUM(N{first}:N{last})")
    m2t.font, m2t.alignment, m2t.number_format = bold, center, "0.00"
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:N{last}"
    return name, r


totals = {}
for grp in model["pjovimas"]:
    totals[grp["grupe"]] = (grp,) + cut_sheet(grp)

# ---- Furnitūra
ws = wb.create_sheet("Furnitūra")
ws["A1"] = f'{model["projektas"]["pavadinimas"]} — furnitūra'
ws["A1"].font = h1
header_row(ws, 3, ["Nr.", "Pavadinimas", "Kiekis", "Vnt.", "Pastaba"], [5, 36, 8, 8, 40])
for i, h in enumerate(model["suvestine"]["furnitura"], 1):
    for c, v in enumerate([i, h["pavadinimas"], h["kiekis"], h["vnt"], h["pastaba"]], 1):
        cell = ws.cell(row=3 + i, column=c, value=v)
        cell.font, cell.border = font, border
        cell.alignment = center if c in (1, 3, 4) else left

# ---- Suvestinė
ws = wb.create_sheet("Suvestinė", 0)
ws["A1"] = f'{model["projektas"]["pavadinimas"]} — medžiagų suvestinė'
ws["A1"].font = h1
ws["A2"] = f'Siena {model["projektas"]["siena"]["ilgis"]} mm, aukštis {model["projektas"]["siena"]["aukstis"]} mm; korpusai: ' + \
    ", ".join(f'{k["id"]} {k["plotis"]}' for k in model["korpusai"])
ws["A2"].font = grey
LP = model["standartas"]["lapas"]
header_row(ws, 4, ["Medžiaga", "Dekoras", "Plotas, m²", f'Lapų ({LP["plotis"]}×{LP["aukstis"]})', "Išeiga"], [38, 14, 12, 18, 9])
ws["F4"] = "Išeigos koeficientas:"; ws["F4"].font = grey
ws["G4"] = LP["iseiga"]; ws["G4"].font = font; ws["G4"].fill = input_fill
ws["F5"] = "Lapo plotas, m²:"; ws["F5"].font = grey
ws["G5"] = f'={LP["plotis"]}*{LP["aukstis"]}/1000000'; ws["G5"].font = font; ws["G5"].number_format = "0.00"
ws.column_dimensions["F"].width = 20
r = 5
for g, (grp, sheet, trow) in totals.items():
    ws.cell(row=r, column=1, value=grp["pavadinimas"]).font = font
    ws.cell(row=r, column=2, value=grp["dekoras"]).font = font
    m2 = ws.cell(row=r, column=3, value=f"='{sheet}'!N{trow}")
    m2.font, m2.number_format, m2.alignment = font, "0.00", center
    if not grp.get("be_lapu"):
        lp = ws.cell(row=r, column=4, value=f"=ROUNDUP(C{r}/($G$5*$G$4),0)")
        lp.font, lp.alignment = font, center
    else:
        ws.cell(row=r, column=4, value="atskiras tiekėjas").font = grey
    for c in range(1, 5):
        ws.cell(row=r, column=c).border = border
    r += 1
r += 1
header_row(ws, r, ["Kantas ABS", "", "Metrų (+5 %)"])
r += 1
for k in model["suvestine"]["kantai"]:
    ws.cell(row=r, column=1, value=f'{k["storis"]} mm').font = font
    c = ws.cell(row=r, column=3, value=k["m"]); c.font, c.alignment = font, center
    for cc in (1, 2, 3):
        ws.cell(row=r, column=cc).border = border
    r += 1
r += 1
ws.cell(row=r, column=1, value="Geltonas langelis — keičiamas (išeigos koeficientas). Lapų skaičius apskaičiuojamas formule iš pjovimo lapų sumų.").font = grey

# ---- Standartas
ws = wb.create_sheet("Standartas")
ws["A1"] = "Konstrukcinis standartas (assets/standartas.json — keičia gamintojas)"
ws["A1"].font = h1
header_row(ws, 3, ["Grupė", "Parametras", "Reikšmė"], [20, 26, 34])
r = 4
for g, params in model["standartas"].items():
    for k, v in params.items():
        for c, val in enumerate([g, k, v], 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font, cell.border = font, border
        r += 1

out = Path(args.out) if args.out else MODEL.with_name("detales.xlsx")
wb.save(out)
print("saved", out)
